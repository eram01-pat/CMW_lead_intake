"""
Excel report for the tender industry classification run.

Produces one workbook with the summary the analysis is for, plus the detail
rows behind it so any number in the summary can be traced back to tenders.
"""

from __future__ import annotations

import logging
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Any, Iterable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.classification.taxonomy import GROUPS, INDUSTRIES, UNCLASSIFIED, group_for

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TOTAL_FONT = Font(bold=True)
_PERCENT_FORMAT = "0.0%"

# Order used for the decision cross-tab columns.
_DECISIONS = ["yes", "maybe", "no"]
_DECISION_LABELS = {
    "yes": "CMW: Yes",
    "maybe": "CMW: Maybe",
    "no": "CMW: No",
}
_NOT_JUDGED = "Not judged"


def _write_sheet(
    ws,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    widths: Optional[Sequence[int]] = None,
    percent_columns: Sequence[int] = (),
    autofilter: bool = True,
    total_row: bool = False,
) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    row_count = 0
    for row in rows:
        ws.append(list(row))
        row_count += 1

    for col_index in percent_columns:
        letter = get_column_letter(col_index)
        for cell in ws[letter][1:]:
            cell.number_format = _PERCENT_FORMAT

    if total_row and row_count:
        for cell in ws[row_count + 1]:
            cell.font = _TOTAL_FONT

    ws.freeze_panes = "A2"
    if autofilter and row_count:
        last_row = row_count + 1 - (1 if total_row else 0)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width


def _share(count: int, total: int) -> float:
    return (count / total) if total else 0.0


def _fmt_date(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value) if value else ""


def _industry_sort_key(item: tuple[str, int]) -> tuple:
    """Most common first; ties broken by taxonomy order for stable output."""
    industry, count = item
    order = INDUSTRIES.index(industry) if industry in INDUSTRIES else len(INDUSTRIES)
    return (-count, order, industry)


def build_workbook(
    rows: Sequence[dict],
    output_path: str,
    *,
    run_metadata: dict[str, Any],
) -> None:
    """
    Write the classification workbook.

    `rows` are the merged tender + classification records produced by
    classify_industries.py — one dict per tender, always carrying `industry`
    and `industry_group`.
    """
    total = len(rows)

    industry_counts: Counter[str] = Counter(r["industry"] for r in rows)
    group_counts: Counter[str] = Counter(r["industry_group"] for r in rows)
    source_group_counts: dict[str, Counter[str]] = defaultdict(Counter)
    industry_decisions: dict[str, Counter[str]] = defaultdict(Counter)
    confidence_counts: Counter[str] = Counter(r.get("confidence") or "unknown" for r in rows)
    secondary_counts: Counter[str] = Counter(
        r["secondary_industry"] for r in rows if r.get("secondary_industry")
    )

    for row in rows:
        source_group_counts[row.get("source_name") or "(unknown)"][row["industry_group"]] += 1
        decision = (row.get("llm_decision") or "").strip().lower()
        industry_decisions[row["industry"]][decision if decision in _DECISIONS else _NOT_JUDGED] += 1

    wb = Workbook()

    # ── Summary by industry — the headline sheet ──────────────────────────
    ws = wb.active
    ws.title = "Summary by Industry"
    ordered_industries = sorted(industry_counts.items(), key=_industry_sort_key)
    summary_rows = [
        [industry, group_for(industry), count, _share(count, total)]
        for industry, count in ordered_industries
    ]
    summary_rows.append(["TOTAL", "", total, _share(total, total)])
    _write_sheet(
        ws,
        ["Industry", "Group", "Tenders", "% of Total"],
        summary_rows,
        widths=[44, 34, 12, 12],
        percent_columns=[4],
        total_row=True,
    )

    # ── Summary by group ─────────────────────────────────────────────────
    ws = wb.create_sheet("Summary by Group")
    ordered_groups = sorted(
        group_counts.items(),
        key=lambda item: (-item[1], GROUPS.index(item[0]) if item[0] in GROUPS else len(GROUPS)),
    )
    group_rows = [[group, count, _share(count, total)] for group, count in ordered_groups]
    group_rows.append(["TOTAL", total, _share(total, total)])
    _write_sheet(
        ws,
        ["Group", "Tenders", "% of Total"],
        group_rows,
        widths=[36, 12, 12],
        percent_columns=[3],
        total_row=True,
    )

    # ── Group by source (cross-tab) ──────────────────────────────────────
    ws = wb.create_sheet("Group by Source")
    present_groups = [g for g in GROUPS if group_counts.get(g)]
    crosstab_rows = []
    for source in sorted(source_group_counts):
        counts = source_group_counts[source]
        crosstab_rows.append(
            [source] + [counts.get(g, 0) for g in present_groups] + [sum(counts.values())]
        )
    crosstab_rows.append(
        ["TOTAL"] + [group_counts.get(g, 0) for g in present_groups] + [total]
    )
    _write_sheet(
        ws,
        ["Source"] + present_groups + ["Total"],
        crosstab_rows,
        widths=[32] + [18] * len(present_groups) + [10],
        total_row=True,
    )

    # ── Industry vs the pipeline's existing CMW relevance decision ───────
    ws = wb.create_sheet("Industry vs CMW Decision")
    decision_rows = []
    for industry, count in ordered_industries:
        counts = industry_decisions[industry]
        decision_rows.append(
            [industry]
            + [counts.get(d, 0) for d in _DECISIONS]
            + [counts.get(_NOT_JUDGED, 0), count]
        )
    decision_rows.append(
        ["TOTAL"]
        + [sum(c.get(d, 0) for c in industry_decisions.values()) for d in _DECISIONS]
        + [
            sum(c.get(_NOT_JUDGED, 0) for c in industry_decisions.values()),
            total,
        ]
    )
    _write_sheet(
        ws,
        ["Industry"] + [_DECISION_LABELS[d] for d in _DECISIONS] + [_NOT_JUDGED, "Total"],
        decision_rows,
        widths=[44, 12, 14, 12, 14, 10],
        total_row=True,
    )

    # ── Confidence + secondary industries ────────────────────────────────
    ws = wb.create_sheet("Confidence")
    confidence_rows = [
        [level, confidence_counts.get(level, 0), _share(confidence_counts.get(level, 0), total)]
        for level in ("high", "medium", "low")
        if confidence_counts.get(level)
    ]
    for level, count in sorted(confidence_counts.items()):
        if level not in ("high", "medium", "low"):
            confidence_rows.append([level, count, _share(count, total)])
    confidence_rows.append(["TOTAL", total, _share(total, total)])
    _write_sheet(
        ws,
        ["Confidence", "Tenders", "% of Total"],
        confidence_rows,
        widths=[18, 12, 12],
        percent_columns=[3],
        total_row=True,
    )

    ws = wb.create_sheet("Secondary Industries")
    _write_sheet(
        ws,
        ["Secondary Industry", "Tenders"],
        [
            [industry, count]
            for industry, count in sorted(secondary_counts.items(), key=_industry_sort_key)
        ]
        or [["(none assigned)", 0]],
        widths=[44, 12],
    )

    # ── Full detail ──────────────────────────────────────────────────────
    ws = wb.create_sheet("All Tenders")
    detail_rows = (
        [
            row.get("source_name") or "",
            row.get("reference_no") or "",
            row.get("title") or "",
            row["industry"],
            row["industry_group"],
            row.get("secondary_industry") or "",
            row.get("confidence") or "",
            row.get("note") or "",
            row.get("status") or "",
            _fmt_date(row.get("closing_date")),
            row.get("llm_decision") or "",
            row.get("detail_url") or "",
        ]
        for row in sorted(
            rows,
            key=lambda r: (r["industry_group"], r["industry"], r.get("source_name") or ""),
        )
    )
    _write_sheet(
        ws,
        [
            "Source", "Reference", "Title", "Industry", "Group",
            "Secondary Industry", "Confidence", "Why", "Status",
            "Closing Date", "CMW Decision", "Detail URL",
        ],
        detail_rows,
        widths=[26, 14, 60, 34, 30, 30, 12, 46, 10, 13, 13, 46],
    )

    # ── Run info ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Run Info")
    _write_sheet(
        ws,
        ["Field", "Value"],
        [[key, str(value)] for key, value in run_metadata.items()],
        widths=[30, 70],
        autofilter=False,
    )

    wb.save(output_path)
    logger.info("Wrote %s (%d tenders, %d industries)", output_path, total, len(industry_counts))


def log_summary(rows: Sequence[dict], top_n: int = 15) -> None:
    """Print the headline counts to the run log / Actions console."""
    total = len(rows)
    counts = Counter(r["industry"] for r in rows)
    logger.info("── Industry summary (%d tenders) ──", total)
    for industry, count in sorted(counts.items(), key=_industry_sort_key)[:top_n]:
        logger.info("  %-46s %5d  (%4.1f%%)", industry, count, 100 * _share(count, total))
    if len(counts) > top_n:
        logger.info("  … and %d more industries (see the Excel report)", len(counts) - top_n)
    unclassified = counts.get(UNCLASSIFIED, 0)
    if unclassified:
        logger.info("  %d tender(s) could not be categorised from the available text", unclassified)


def write_markdown_summary(rows: Sequence[dict], path: str, *, run_metadata: dict[str, Any]) -> None:
    """
    Write the same headline numbers as Markdown.

    The classification workflow pipes this into the GitHub Actions run summary
    so the counts are readable without downloading the workbook.
    """
    total = len(rows)
    industry_counts: Counter[str] = Counter(r["industry"] for r in rows)
    group_counts: Counter[str] = Counter(r["industry_group"] for r in rows)

    lines = [
        "## Tender industry classification",
        "",
        f"**{total:,} tenders** classified into **{len(industry_counts)} industries**.",
        "",
        "### By group",
        "",
        "| Group | Tenders | % of total |",
        "| --- | ---: | ---: |",
    ]
    for group, count in sorted(
        group_counts.items(),
        key=lambda item: (-item[1], GROUPS.index(item[0]) if item[0] in GROUPS else len(GROUPS)),
    ):
        lines.append(f"| {group} | {count:,} | {100 * _share(count, total):.1f}% |")

    lines += [
        "",
        "### By industry",
        "",
        "| Industry | Group | Tenders | % of total |",
        "| --- | --- | ---: | ---: |",
    ]
    for industry, count in sorted(industry_counts.items(), key=_industry_sort_key):
        lines.append(
            f"| {industry} | {group_for(industry)} | {count:,} | "
            f"{100 * _share(count, total):.1f}% |"
        )

    lines += ["", "### Run details", "", "| Field | Value |", "| --- | --- |"]
    for key, value in run_metadata.items():
        lines.append(f"| {key} | {value} |")
    lines.append("")

    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    logger.info("Wrote %s", path)
